# Genera landing/plantilla-packing-list-acopiar.xlsx (imán para importadores). Uso: .venv/bin/python scripts/plantilla_packing_list.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb = Workbook(); ws = wb.active; ws.title = 'Packing List'
verde = PatternFill('solid', fgColor='0F766E'); gris = PatternFill('solid', fgColor='EEF2EE'); blanco = Font(color='FFFFFF', bold=True); b = Font(bold=True)
fino = Side(style='thin', color='D9E0DB'); borde = Border(left=fino, right=fino, top=fino, bottom=fino)
ws.merge_cells('A1:J1'); ws['A1'] = 'PACKING LIST / LISTA DE EMPAQUE'; ws['A1'].font = Font(size=16, bold=True, color='0F766E')
ws.merge_cells('A2:J2'); ws['A2'] = 'Plantilla gratuita de Acopiar · acopiar.app · Completa el encabezado y una fila por producto dentro de cada caja o pallet.'; ws['A2'].font = Font(italic=True, color='5E6B66')
cab = [('Exportador / Vendedor', 'B4'), ('Consignatario / Comprador', 'B5'), ('N° de factura (Invoice)', 'B6'), ('Fecha', 'B7'), ('Lugar de carga', 'G4'), ('Lugar de destino', 'G5'), ('Modo (marítimo / aéreo / terrestre)', 'G6'), ('N° de contenedor / guía', 'G7')]
for et, celda in cab:
    col = celda[0]; fila = int(celda[1:]); etcol = chr(ord(col) - 1)
    ws[f'{etcol}{fila}'] = et; ws[f'{etcol}{fila}'].font = b; ws[celda].fill = gris; ws[celda].border = borde
    ws.merge_cells(f'{celda}:{chr(ord(col)+2)}{fila}')
hdr = ['Pallet', 'Caja', 'Descripción del producto', 'Código HS', 'Cantidad', 'Unidad / presentación', 'Peso neto (kg)', 'Peso bruto (kg)', 'Dimensiones caja (cm)', 'Valor unitario (USD)']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=9, column=i, value=h); c.fill = verde; c.font = blanco; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = borde
ejemplo = [['PAL-001', 'BQ-0001', 'Paracetamol 500 mg x16 comprimidos', '3004.90', 24, 'caja de 16', 1.2, 1.5, '40x30x30', 2.10], ['PAL-001', 'BQ-0002', 'Toallitas húmedas x80', '3401.11', 10, 'paquete', 6.0, 6.5, '40x30x30', 1.80]]
for r, fila in enumerate(ejemplo, 10):
    for i, v in enumerate(fila, 1):
        c = ws.cell(row=r, column=i, value=v); c.border = borde
for r in range(12, 60):
    for i in range(1, 11): ws.cell(row=r, column=i).border = borde
ws['A61'] = 'TOTALES'; ws['A61'].font = b
ws['E61'] = '=SUM(E10:E60)'; ws['G61'] = '=SUM(G10:G60)'; ws['H61'] = '=SUM(H10:H60)'; ws['J61'] = '=SUMPRODUCT(E10:E60,J10:J60)'
for c in ['E61','G61','H61','J61']: ws[c].font = b; ws[c].fill = gris
ws['A63'] = 'Notas: el código HS es referencial, confírmalo con tu agente de aduana. Peso bruto = neto + embalaje. Una fila por producto; repite el número de caja si una caja trae varios productos.'
ws['A63'].font = Font(italic=True, color='5E6B66'); ws.merge_cells('A63:J63')
for i, w in enumerate([10, 10, 38, 11, 10, 20, 14, 14, 20, 16], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[9].height = 32; ws.freeze_panes = 'A10'
ws2 = wb.create_sheet('Cómo usarla')
textos = ['CÓMO USAR ESTA PLANTILLA', '', '1. Completa el encabezado: quién envía, quién recibe, número de factura, fecha, origen y destino.', '2. Una fila por producto. Si una caja trae varios productos, repite el número de caja en cada fila.', '3. Agrupa las cajas por pallet (columna Pallet). Si despachas sin pallets, deja la columna vacía.', '4. Código HS: es el código arancelario. Si no lo sabes, pregúntale a tu agente de aduana o busca "clasificación arancelaria + producto".', '5. Pesos: neto = solo el producto; bruto = producto + caja + embalaje. Aduana y transportista usan el bruto.', '6. Los totales se calculan solos en la fila 61.', '', '¿Cansado de armar esto a mano para cada despacho?', 'Acopiar lo genera solo: registras cada caja con su QR desde el celular, armas el pallet escaneando y el packing list sale listo al cerrar el despacho. Gratis para empezar: https://acopiar.app']
for r, t in enumerate(textos, 1):
    ws2.cell(row=r, column=1, value=t).font = Font(bold=(r in (1, 10)), size=14 if r == 1 else 11)
ws2.column_dimensions['A'].width = 120
wb.save('landing/plantilla-packing-list-acopiar.xlsx'); print('plantilla ok')
